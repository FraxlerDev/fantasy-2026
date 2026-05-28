from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Inches, Pt, RGBColor


OUTPUT = "Fantasy_2026_UA_MVP_Product_Spec.docx"


ACCENT = RGBColor(20, 88, 62)
MUTED = RGBColor(88, 98, 110)
LIGHT = "EAF3EE"
HEADER = "D9EADF"
BORDER = "B8C7BE"


def set_cell_shading(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = tc_pr.find(qn("w:shd"))
    if shd is None:
        shd = OxmlElement("w:shd")
        tc_pr.append(shd)
    shd.set(qn("w:fill"), fill)


def set_cell_text(cell, text, bold=False, color=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = p.add_run(text)
    r.bold = bold
    r.font.name = "Arial"
    r.font.size = Pt(9.5)
    if color:
        r.font.color.rgb = color
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER


def set_table_borders(table, color=BORDER, size="6"):
    tbl = table._tbl
    tbl_pr = tbl.tblPr
    borders = tbl_pr.first_child_found_in("w:tblBorders")
    if borders is None:
        borders = OxmlElement("w:tblBorders")
        tbl_pr.append(borders)
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        tag = f"w:{edge}"
        element = borders.find(qn(tag))
        if element is None:
            element = OxmlElement(tag)
            borders.append(element)
        element.set(qn("w:val"), "single")
        element.set(qn("w:sz"), size)
        element.set(qn("w:space"), "0")
        element.set(qn("w:color"), color)


def set_cell_margins(table, top=90, start=120, bottom=90, end=120):
    tbl_pr = table._tbl.tblPr
    margins = tbl_pr.first_child_found_in("w:tblCellMar")
    if margins is None:
        margins = OxmlElement("w:tblCellMar")
        tbl_pr.append(margins)
    for m, v in (("top", top), ("start", start), ("bottom", bottom), ("end", end)):
        node = margins.find(qn(f"w:{m}"))
        if node is None:
            node = OxmlElement(f"w:{m}")
            margins.append(node)
        node.set(qn("w:w"), str(v))
        node.set(qn("w:type"), "dxa")


def style_table(table):
    table.alignment = WD_TABLE_ALIGNMENT.LEFT
    table.style = "Table Grid"
    table.autofit = False
    set_table_borders(table)
    set_cell_margins(table)
    for row_idx, row in enumerate(table.rows):
        for cell in row.cells:
            if row_idx == 0:
                set_cell_shading(cell, HEADER)
                for paragraph in cell.paragraphs:
                    for run in paragraph.runs:
                        run.bold = True
                        run.font.color.rgb = RGBColor(24, 50, 38)


def set_column_widths(table, widths):
    for row in table.rows:
        for idx, width in enumerate(widths):
            row.cells[idx].width = width


def add_table(doc, headers, rows, widths=None):
    table = doc.add_table(rows=1, cols=len(headers))
    hdr = table.rows[0].cells
    for idx, header in enumerate(headers):
        set_cell_text(hdr[idx], header, bold=True, color=RGBColor(24, 50, 38))
    for row in rows:
        cells = table.add_row().cells
        for idx, value in enumerate(row):
            set_cell_text(cells[idx], str(value))
    if widths:
        set_column_widths(table, widths)
    style_table(table)
    return table


def add_bullets(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Bullet")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_numbered(doc, items):
    for item in items:
        p = doc.add_paragraph(style="List Number")
        p.paragraph_format.space_after = Pt(3)
        p.add_run(item)


def add_callout(doc, title, body):
    table = doc.add_table(rows=1, cols=1)
    cell = table.rows[0].cells[0]
    set_cell_shading(cell, LIGHT)
    cell.text = ""
    p = cell.paragraphs[0]
    r = p.add_run(title)
    r.bold = True
    r.font.name = "Arial"
    r.font.size = Pt(10.5)
    r.font.color.rgb = ACCENT
    p2 = cell.add_paragraph()
    r2 = p2.add_run(body)
    r2.font.name = "Arial"
    r2.font.size = Pt(9.5)
    set_table_borders(table, color="C7D8CF")
    set_cell_margins(table, top=120, start=160, bottom=120, end=160)


def configure_document(doc):
    section = doc.sections[0]
    section.top_margin = Inches(0.65)
    section.bottom_margin = Inches(0.65)
    section.left_margin = Inches(0.5)
    section.right_margin = Inches(0.5)

    styles = doc.styles
    normal = styles["Normal"]
    normal.font.name = "Arial"
    normal.font.size = Pt(10.5)
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.line_spacing = 1.08

    for style_name, size, color in [
        ("Title", 24, ACCENT),
        ("Heading 1", 16, ACCENT),
        ("Heading 2", 13, RGBColor(32, 55, 45)),
        ("Heading 3", 11.5, RGBColor(32, 55, 45)),
    ]:
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(size)
        style.font.bold = True
        style.font.color.rgb = color
        style.paragraph_format.space_before = Pt(10 if style_name != "Title" else 0)
        style.paragraph_format.space_after = Pt(5)

    for style_name in ("List Bullet", "List Number"):
        style = styles[style_name]
        style.font.name = "Arial"
        style.font.size = Pt(10)
        style.paragraph_format.space_after = Pt(3)

    header = section.header
    hp = header.paragraphs[0]
    hp.alignment = WD_ALIGN_PARAGRAPH.LEFT
    hr = hp.add_run("Fantasy 2026 UA | MVP Product Spec")
    hr.font.name = "Arial"
    hr.font.size = Pt(8.5)
    hr.font.color.rgb = MUTED

    footer = section.footer
    fp = footer.paragraphs[0]
    fp.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    fr = fp.add_run("Сторінка ")
    fr.font.name = "Arial"
    fr.font.size = Pt(8.5)
    field = OxmlElement("w:fldSimple")
    field.set(qn("w:instr"), "PAGE")
    fp._p.append(field)


def build_doc():
    doc = Document()
    configure_document(doc)

    title = doc.add_paragraph(style="Title")
    title.add_run("MVP Product Spec: Fantasy 2026 UA")
    subtitle = doc.add_paragraph()
    subtitle.alignment = WD_ALIGN_PARAGRAPH.LEFT
    r = subtitle.add_run(
        "Production-oriented специфікація першого релізу українського fantasy-футболу до ЧС-2026"
    )
    r.font.name = "Arial"
    r.font.size = Pt(11)
    r.font.color.rgb = MUTED

    add_table(
        doc,
        ["Параметр", "Рішення"],
        [
            ["Аудиторія", "Українські футбольні фанати, спільнота друзів і приватних ліг"],
            ["Мова", "Українська"],
            ["Платформа", "Desktop-first web"],
            ["Модель", "Free-to-play, без entry fee і cash-prize у MVP"],
            ["Режим", "Одна команда користувача на весь ЧС-2026"],
            ["Дані", "Ручне введення через admin-панель після зіграного матчу"],
            ["Стек", "Next.js + NestJS + PostgreSQL + Redis"],
        ],
        [Inches(1.7), Inches(5.8)],
    )

    doc.add_heading("1. Executive Summary", level=1)
    doc.add_paragraph(
        "Fantasy 2026 UA — це український free-to-play fantasy-продукт до ЧС-2026, "
        "побудований навколо однієї команди користувача на весь турнір, приватних ліг "
        "для друзів, зрозумілого balanced scoring і медійного onboarding-шару. MVP "
        "не використовує live-data provider: статистика матчів вводиться адміністраторами "
        "вручну після матчу, після чого система автоматично перераховує очки, турові "
        "результати та leaderboard-и."
    )
    add_callout(
        doc,
        "Головний принцип MVP",
        "Перший реліз має бути production-придатним, але обмеженим за поверхнею: стабільні правила, "
        "надійна адмінка для ручних даних, приватні ліги та прозорий перерахунок очок важливіші за live-ефекти."
    )

    doc.add_heading("2. Product Goals", level=1)
    add_table(
        doc,
        ["Ціль", "Опис", "Критерій успіху"],
        [
            [
                "Швидкий старт",
                "Користувач має зрозуміти гру і створити команду без знання fantasy-термінів.",
                "Створення команди за 2-4 хвилини.",
            ],
            [
                "Соціальна гра",
                "Основна мотивація — змагатися з друзями у приватних лігах.",
                "Користувач створює або вступає в лігу після команди.",
            ],
            [
                "Прозорі очки",
                "Balanced scoring має бути легким для пояснення і перевірки.",
                "Кожен користувач бачить деталізацію очок по гравцю.",
            ],
            [
                "Операційний контроль",
                "Адміни можуть вручну вести турнір без data provider.",
                "Матч можна закрити, перерахувати і зааудитити з backoffice.",
            ],
        ],
        [Inches(1.45), Inches(3.55), Inches(2.5)],
    )

    doc.add_heading("3. MVP Scope", level=1)
    doc.add_heading("3.1 In Scope", level=2)
    add_bullets(
        doc,
        [
            "Реєстрація та вхід користувача.",
            "Створення однієї fantasy-команди на весь турнір.",
            "Каталог збірних, гравців, позицій і цін.",
            "Squad builder із бюджетом, позиційними лімітами та лімітом гравців зі збірної.",
            "Вибір стартових 11, лавки, капітана і віцекапітанa.",
            "Трансфери між matchday/stage з безкоштовним лімітом і штрафами.",
            "Balanced scoring після ручного введення статистики матчу.",
            "Приватні ліги, invite code, вступ до ліги та leaderboard.",
            "Глобальний leaderboard.",
            "Медійні сторінки: правила, швидкий старт, гайди, поради.",
            "Admin-панель для даних, матчів, статистики, перерахунку, скарг і fraud-позначок.",
        ],
    )

    doc.add_heading("3.2 Out of Scope", level=2)
    add_bullets(
        doc,
        [
            "Live-оновлення під час матчу.",
            "Інтеграція з реальним футбольним data provider у MVP.",
            "Мобільні застосунки iOS/Android.",
            "Paid contests, entry fee, cash-prize або betting-функціональність.",
            "Коментарі, чати, activity feed і reaction system.",
            "Advanced scoring на базі defensive contributions або event-derived bonus.",
            "Мультимовність.",
        ],
    )

    doc.add_heading("4. Game Rules", level=1)
    add_table(
        doc,
        ["Елемент", "MVP-правило"],
        [
            ["Склад", "15 гравців: 2 GK, 5 DEF, 5 MID, 3 FWD"],
            ["Стартовий склад", "11 гравців у дозволеній формації"],
            ["Бюджет", "100 кредитів"],
            ["Капітан", "Подвоює очки обраного гравця"],
            ["Віцекапітан", "Отримує x2, якщо капітан зіграв 0 хвилин"],
            ["Дедлайн", "До старту першого матчу поточного matchday/stage"],
            ["Одна команда", "Один користувач може мати одну основну fantasy-команду"],
            ["Після дедлайну", "Склад, капітан і трансфери для stage заблоковані"],
        ],
        [Inches(2.2), Inches(5.3)],
    )

    doc.add_heading("4.1 Stage Model", level=2)
    add_table(
        doc,
        ["Stage", "Free transfers", "Max гравців з однієї збірної"],
        [
            ["Pre-tournament", "Unlimited", "2"],
            ["Group MD2", "2", "2"],
            ["Group MD3", "2", "2"],
            ["Round of 32", "4", "3"],
            ["Round of 16", "4", "4"],
            ["Quarter-finals", "5", "4"],
            ["Semi-finals", "5", "5"],
            ["Finals", "6", "8"],
        ],
        [Inches(2.7), Inches(2.2), Inches(2.6)],
    )
    doc.add_paragraph(
        "Кожен додатковий трансфер понад ліміт коштує -4 очки. Точні назви stage можна "
        "адаптувати під офіційний календар після затвердження матчів."
    )

    doc.add_heading("5. Balanced Scoring", level=1)
    add_table(
        doc,
        ["Подія", "Очки"],
        [
            ["Вихід на поле 1-59 хв", "+1"],
            ["60+ хвилин", "+2"],
            ["Гол GK/DEF", "+10"],
            ["Гол MID", "+6"],
            ["Гол FWD", "+4"],
            ["Асист", "+3"],
            ["Clean sheet GK/DEF", "+4"],
            ["Clean sheet MID", "+1"],
            ["Кожні 2 пропущені голи GK/DEF", "-1"],
            ["Кожні 3 сейви GK", "+1"],
            ["Відбитий пенальті", "+5"],
            ["Нереалізований пенальті", "-2"],
            ["Жовта картка", "-1"],
            ["Червона картка", "-3"],
            ["Автогол", "-2"],
            ["Player of the Match", "+3"],
        ],
        [Inches(5.5), Inches(2.0)],
    )
    add_callout(
        doc,
        "Ruleset engine",
        "Scoring не слід хардкодити в UI. Правила мають зберігатися як versioned ruleset, щоб "
        "адмін міг бачити активну версію, а розробники могли тестувати зміни без переписування доменної логіки."
    )

    doc.add_heading("6. User Roles", level=1)
    add_table(
        doc,
        ["Роль", "Права"],
        [
            ["Guest", "Читає landing, правила, гайди; може почати onboarding."],
            ["User", "Створює команду, змінює склад до дедлайну, вступає в ліги, дивиться очки."],
            ["League owner", "Створює приватну лігу, поширює invite code, бачить учасників."],
            ["Admin", "Керує турнірними даними, гравцями, матчами, статистикою і перерахунком."],
            ["Super admin", "Керує ролями, ручними корекціями, fraud flags і системними налаштуваннями."],
        ],
        [Inches(1.65), Inches(5.85)],
    )

    doc.add_heading("7. User Experience", level=1)
    doc.add_heading("7.1 Core User Flow", level=2)
    add_numbered(
        doc,
        [
            "Користувач заходить на головну сторінку і бачить просте пояснення гри.",
            "Натискає 'Швидкий старт' і проходить короткий onboarding.",
            "Реєструється або входить.",
            "Створює fantasy-команду з 15 гравців у межах бюджету.",
            "Вибирає стартові 11, капітана і віцекапітанa.",
            "Створює приватну лігу або приєднується за invite code.",
            "Після матчів переглядає деталізацію очок і leaderboard.",
            "Перед наступним stage робить трансфери.",
        ],
    )

    doc.add_heading("7.2 Screens", level=2)
    add_table(
        doc,
        ["Екран", "Призначення", "MVP-вимоги"],
        [
            [
                "Landing",
                "Пояснити гру і привести до створення команди.",
                "Hero, короткі правила, CTA, блоки гайдів, free-to-play сигнал.",
            ],
            [
                "Швидкий старт",
                "Навчити новачка за 2 хвилини.",
                "3-5 кроків: склад, бюджет, капітан, ліги, дедлайни.",
            ],
            [
                "Squad Builder",
                "Зібрати 15 гравців.",
                "Фільтри за збірною/позицією/ціною, бюджет, warnings, пошук.",
            ],
            [
                "Lineup",
                "Налаштувати стартовий склад.",
                "Поле/схема, лавка, captain/vice, deadline state.",
            ],
            [
                "Transfers",
                "Змінити склад між stage.",
                "Free transfers, hit cost, budget after changes, eliminated players.",
            ],
            [
                "League Hub",
                "Грати з друзями.",
                "Invite code, список учасників, leaderboard, позиція користувача.",
            ],
            [
                "Player Detail",
                "Допомогти вибору гравця.",
                "Ціна, позиція, збірна, матчі, очки, ownership пізніше.",
            ],
            [
                "Guides",
                "Медійний шар продукту.",
                "Правила, поради, captain picks, пояснення scoring.",
            ],
            [
                "Profile",
                "Базові налаштування.",
                "Username, email, команда, ліги, opt-ins.",
            ],
        ],
        [Inches(1.55), Inches(2.25), Inches(3.7)],
    )

    doc.add_heading("8. Admin Product Spec", level=1)
    add_table(
        doc,
        ["Admin-модуль", "Функції"],
        [
            [
                "Teams",
                "Створення/редагування збірних, груп, прапорів/кольорів, статусу участі.",
            ],
            [
                "Players",
                "Гравці, позиції, ціни, збірна, availability status, notes.",
            ],
            [
                "Fixtures",
                "Матчі, stage, kickoff, статус: scheduled, locked, played, settled.",
            ],
            [
                "Match Stats",
                "Введення хвилин, голів, асистів, карток, сейвів, пенальті, clean sheet, POTM.",
            ],
            [
                "Scoring",
                "Запуск перерахунку, перегляд score entries, ручна корекція з причиною.",
            ],
            [
                "Users & Teams",
                "Пошук користувачів, перегляд команд, locks, fraud flags.",
            ],
            [
                "Leagues",
                "Перегляд приватних ліг, учасників, invite code, скарг.",
            ],
            [
                "Complaints/Fraud",
                "Список скарг, статус review, fraud-позначки, нотатки модератора.",
            ],
            [
                "Audit Log",
                "Хто, коли і що змінив у критичних даних.",
            ],
        ],
        [Inches(1.75), Inches(5.75)],
    )
    add_callout(
        doc,
        "Admin acceptance rule",
        "Жодна ручна корекція очок або статистики не може бути без reason, actor, timestamp і audit trail."
    )

    doc.add_heading("9. Data Model", level=1)
    add_table(
        doc,
        ["Сутність", "Ключові поля"],
        [
            ["users", "id, email, username, role, locale, created_at, status"],
            ["fantasy_teams", "id, user_id, name, budget, total_points, created_at"],
            ["national_teams", "id, name_uk, group_key, status"],
            ["players", "id, national_team_id, name, position, price, status"],
            ["stages", "id, key, name, deadline_at, free_transfers, nation_limit"],
            ["fixtures", "id, stage_id, home_team_id, away_team_id, kickoff_at, status"],
            ["roster_entries", "fantasy_team_id, player_id, stage_id, purchase_price"],
            ["lineup_versions", "fantasy_team_id, stage_id, formation, captain_id, vice_id, saved_at"],
            ["transfers", "fantasy_team_id, stage_id, out_player_id, in_player_id, hit_cost"],
            ["match_player_stats", "fixture_id, player_id, minutes, goals, assists, cards, saves, potm"],
            ["score_entries", "fantasy_team_id, fixture_id, player_id, points, breakdown_json, status"],
            ["leagues", "id, owner_user_id, name, invite_code, visibility"],
            ["league_members", "league_id, fantasy_team_id, joined_at"],
            ["complaints", "id, reporter_user_id, entity_type, entity_id, reason, status"],
            ["audit_logs", "actor_id, action, entity_type, entity_id, before_json, after_json, created_at"],
        ],
        [Inches(2.15), Inches(5.35)],
    )

    doc.add_heading("10. API Outline", level=1)
    add_table(
        doc,
        ["Метод", "Endpoint", "Призначення"],
        [
            ["GET", "/v1/content/home", "Дані для головної та медійних блоків"],
            ["GET", "/v1/tournament/stages", "Список stage, дедлайни, правила"],
            ["GET", "/v1/players", "Каталог гравців з фільтрами"],
            ["POST", "/v1/fantasy-team", "Створення fantasy-команди"],
            ["GET", "/v1/fantasy-team/me", "Команда поточного користувача"],
            ["POST", "/v1/fantasy-team/lineup", "Збереження стартового складу"],
            ["POST", "/v1/fantasy-team/transfers", "Застосування transfer set"],
            ["POST", "/v1/leagues", "Створення приватної ліги"],
            ["POST", "/v1/leagues/join", "Вступ за invite code"],
            ["GET", "/v1/leagues/{id}/leaderboard", "Leaderboard приватної ліги"],
            ["GET", "/v1/leaderboard/global", "Глобальний leaderboard"],
            ["GET", "/v1/scores/me", "Деталізація очок користувача"],
            ["POST", "/admin/fixtures/{id}/stats", "Введення статистики матчу"],
            ["POST", "/admin/fixtures/{id}/settle", "Перерахунок і settlement матчу"],
            ["POST", "/admin/score-adjustments", "Ручна корекція очок"],
        ],
        [Inches(0.75), Inches(2.75), Inches(4.0)],
    )

    doc.add_heading("11. System Architecture", level=1)
    add_table(
        doc,
        ["Компонент", "Роль у MVP"],
        [
            ["Next.js App Router", "Desktop-first web, SSR для медійних сторінок, authenticated app screens."],
            ["NestJS API", "Auth, fantasy domain, leagues, admin API, scoring orchestration."],
            ["PostgreSQL", "Source of truth для користувачів, складів, матчів, очок і audit logs."],
            ["Redis", "Кеш leaderboard, locks, jobs для перерахунку і rate limiting."],
            ["Worker", "Асинхронний scoring recalculation після admin settlement."],
            ["Admin zone", "Окремий protected route з role-based access control."],
        ],
        [Inches(2.0), Inches(5.5)],
    )

    doc.add_heading("12. Acceptance Criteria", level=1)
    add_table(
        doc,
        ["Функція", "Критерій приймання"],
        [
            ["Створення команди", "Користувач не може зберегти склад, що порушує бюджет, позиції або nation limit."],
            ["Дедлайн", "Після дедлайну stage API відхиляє зміни складу, капітана і трансферів."],
            ["Scoring", "Після settlement матчу очки перераховуються однаково для всіх команд з відповідними гравцями."],
            ["Captain", "Капітан множить очки x2; віцекапітан підхоплює тільки якщо капітан має 0 хвилин."],
            ["Ліги", "Користувач може створити приватну лігу, запросити за кодом і бачити leaderboard."],
            ["Admin stats", "Адмін може ввести статистику матчу, зберегти, запустити settlement і побачити audit log."],
            ["Manual correction", "Корекція очок вимагає reason і відображається в audit trail."],
            ["Мова", "Увесь користувацький інтерфейс MVP українською."],
        ],
        [Inches(1.85), Inches(5.65)],
    )

    doc.add_heading("13. Roadmap", level=1)
    add_table(
        doc,
        ["Етап", "Зміст", "Орієнтир"],
        [
            ["1. Product & UX spec", "Фіналізація правил, IA, wireframes, admin flows.", "1-2 тижні"],
            ["2. Platform foundation", "Repo, CI, auth, DB schema, roles, base layout.", "2 тижні"],
            ["3. Core gameplay", "Squad builder, lineup, captain, transfers, validation.", "3-4 тижні"],
            ["4. Leagues", "Private leagues, invite code, leaderboard.", "2 тижні"],
            ["5. Admin & scoring", "Match stats entry, settlement worker, corrections, audit.", "3-4 тижні"],
            ["6. Content layer", "Landing, quick start, guides, editorial pages.", "1-2 тижні"],
            ["7. QA hardening", "Unit/domain tests, E2E flows, load smoke, security checks.", "2 тижні"],
            ["8. Beta", "Closed beta, bugfixes, scoring validation, ops rehearsal.", "2 тижні"],
        ],
        [Inches(1.75), Inches(4.2), Inches(1.55)],
    )

    doc.add_heading("14. Key Risks", level=1)
    add_table(
        doc,
        ["Ризик", "Вплив", "Мітигація"],
        [
            ["Помилки ручного введення статистики", "Неправильні очки і спори", "Audit log, preview settlement, second admin review."],
            ["Перевантаження правил для новачків", "Низька activation", "Швидкий старт, templates, прості пояснення scoring."],
            ["Суперечки щодо призів/чесності", "Репутаційний ризик", "Free-to-play, прозорі правила, fraud flags, ToS."],
            ["Пізня адмінка", "Неможливо вести турнір", "Admin і scoring будувати до публічної beta."],
            ["Проблеми дедлайнів", "Неконсистентні склади", "Server-side locks, transactional validation, tests."],
        ],
        [Inches(2.1), Inches(2.15), Inches(3.25)],
    )

    doc.add_heading("15. Next Decisions", level=1)
    add_bullets(
        doc,
        [
            "Затвердити назву або лишити Fantasy 2026 UA як робочий бренд.",
            "Підтвердити точні трансферні ліміти по stage.",
            "Підтвердити всі значення balanced scoring перед імплементацією.",
            "Вирішити auth-модель: password, magic link або social login.",
            "Підготувати wireframes для основних екранів і admin-панелі.",
        ],
    )

    doc.save(OUTPUT)


if __name__ == "__main__":
    build_doc()
